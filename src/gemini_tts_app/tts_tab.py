# file-path: src/gemini_tts_app/tts_tab.py
# version: 1.0
# last-updated: 2025-07-22
# description: Module chuyên trách cho tab Text-to-Speech.

import tkinter as tk
from tkinter import ttk, scrolledtext, filedialog, messagebox
import os
import platform
import subprocess
import datetime
import threading
import time
import logging

from .tts_logic import generate_tts_audio_multithreaded
from .constants import (
    DEFAULT_VOICE, MIN_TEMPERATURE, MAX_TEMPERATURE,
    PREDEFINED_READING_STYLES, GEMINI_TTS_VOICES_DETAILED,
    DEFAULT_TEMPERATURE, DEFAULT_TOP_P, MIN_TOP_P, MAX_TOP_P,
    NUM_API_KEYS, COLOR_OK, COLOR_NORMAL, COLOR_WARN, COLOR_ERROR
)
from .utils import get_resource_path

try:
    import docx
except ImportError:
    docx = None

class TTSTab(ttk.Frame):
    def __init__(self, parent, main_app_instance):
        super().__init__(parent, padding="10")
        self.main_app = main_app_instance # Tham chiếu đến instance chính của ứng dụng
        self.settings = self.main_app.settings

        self.voice_display_list = [f"{v['name']} ({v['gender']}) - {v['description']}" for v in GEMINI_TTS_VOICES_DETAILED]
        self.voice_name_list = [v['name'] for v in GEMINI_TTS_VOICES_DETAILED]

        # --- KHỞI TẠO CÁC BIẾN (VARIABLES) ---
        self.selected_voice_name = tk.StringVar(value=self.settings.get("default_voice", DEFAULT_VOICE))
        self.selected_voice_display = tk.StringVar()
        self.temperature_var = tk.DoubleVar(value=self.settings.get("temperature", DEFAULT_TEMPERATURE))
        self.top_p_var = tk.DoubleVar(value=self.settings.get("top_p", DEFAULT_TOP_P))
        self.words_per_chunk_var = tk.IntVar(value=self.settings.get("max_words_per_part", 1000))
        self.output_dir_var = tk.StringVar(value=self.settings.get("save_dir", os.path.expanduser("~")))
        self.story_name_var = tk.StringVar(value="MyStory")

        self.languages = {"Tiếng Việt": "vi", "English": "en"}
        self.selected_language = tk.StringVar(value="Tiếng Việt")

        default_lang_key = self.languages[self.selected_language.get()]
        default_style = PREDEFINED_READING_STYLES.get(default_lang_key, [""])[0]
        self.reading_style_prompt_var = tk.StringVar(value=default_style)

        self.last_saved_output_dir = None
        self.thread_status_labels = []

        self._create_widgets()
        self.update_voice_display(self.selected_voice_name.get())
        self.update_word_count()
        self._on_language_change()

    def _create_widgets(self):
        self.rowconfigure(0, weight=3); self.rowconfigure(1, weight=0); self.rowconfigure(2, weight=0);
        self.rowconfigure(3, weight=0); self.rowconfigure(4, weight=0); self.rowconfigure(5, weight=3)
        self.columnconfigure(0, weight=1)

        input_text_frame = ttk.LabelFrame(self, text="Input Text", padding="10")
        input_text_frame.grid(row=0, column=0, sticky="ewns", padx=5, pady=5)
        input_text_frame.rowconfigure(0, weight=1); input_text_frame.columnconfigure(0, weight=1)
        self.main_text_input = scrolledtext.ScrolledText(input_text_frame, wrap=tk.WORD, height=10)
        self.main_text_input.grid(row=0, column=0, sticky="nsew")
        self.main_text_input.bind("<KeyRelease>", self.update_word_count)

        count_frame = ttk.Frame(input_text_frame)
        count_frame.grid(row=1, column=0, sticky="ew", pady=(5,0), padx=5)
        self.main_word_count_label = ttk.Label(count_frame, text="Ký tự: 0 | Từ: 0 | Dự kiến chunks: 0")
        self.main_word_count_label.pack(side="left")
        self.import_button = ttk.Button(count_frame, text="Import từ File...", command=self.import_text_from_file)
        self.import_button.pack(side="right")

        reading_style_frame = ttk.LabelFrame(self, text="Reading Style Prompt (Select or Type Custom)", padding="10")
        reading_style_frame.grid(row=1, column=0, sticky="ew", padx=5, pady=5)
        self.reading_style_combobox = ttk.Combobox(reading_style_frame, textvariable=self.reading_style_prompt_var, height=10)
        self.reading_style_combobox.pack(fill="x", expand=True, padx=5, pady=5)

        settings_container_frame = ttk.Frame(self)
        settings_container_frame.grid(row=2, column=0, sticky="ew", padx=0, pady=0)
        settings_container_frame.columnconfigure(0, weight=1); settings_container_frame.columnconfigure(1, weight=1)

        gen_settings_frame = ttk.LabelFrame(settings_container_frame, text="Generation Settings", padding="10")
        gen_settings_frame.grid(row=0, column=0, sticky="nsew", padx=5, pady=5)
        gen_settings_frame.columnconfigure(1, weight=1)

        ttk.Label(gen_settings_frame, text="Ngôn ngữ:").grid(row=0, column=0, sticky="w", pady=2)
        self.lang_combobox = ttk.Combobox(gen_settings_frame, textvariable=self.selected_language, values=list(self.languages.keys()), state="readonly")
        self.lang_combobox.grid(row=0, column=1, columnspan=2, sticky="ew", pady=2)
        self.lang_combobox.bind("<<ComboboxSelected>>", self._on_language_change)

        ttk.Label(gen_settings_frame, text="Select Voice:").grid(row=1, column=0, sticky="w", pady=2)
        self.voice_dropdown = ttk.Combobox(gen_settings_frame, textvariable=self.selected_voice_display, values=self.voice_display_list, state="readonly")
        self.voice_dropdown.grid(row=1, column=1, columnspan=2, sticky="ew", pady=2)
        self.voice_dropdown.bind('<<ComboboxSelected>>', self.on_voice_selected)

        ttk.Label(gen_settings_frame, text="Temperature:").grid(row=2, column=0, sticky="w", pady=2)
        self.temp_scale = ttk.Scale(gen_settings_frame, from_=MIN_TEMPERATURE, to=MAX_TEMPERATURE, variable=self.temperature_var, orient=tk.HORIZONTAL, command=lambda v: self.temp_scale_val_label.config(text=f"{float(v):.2f}"))
        self.temp_scale.grid(row=2, column=1, sticky="ew", pady=2)
        self.temp_scale_val_label = ttk.Label(gen_settings_frame, text=f"{self.temperature_var.get():.2f}", width=4)
        self.temp_scale_val_label.grid(row=2, column=2, sticky="w", padx=(5,0))

        ttk.Label(gen_settings_frame, text="Top P:").grid(row=3, column=0, sticky="w", pady=2)
        self.top_p_scale = ttk.Scale(gen_settings_frame, from_=MIN_TOP_P, to=MAX_TOP_P, variable=self.top_p_var, orient=tk.HORIZONTAL, command=lambda v: self.top_p_scale_val_label.config(text=f"{float(v):.2f}"))
        self.top_p_scale.grid(row=3, column=1, sticky="ew", pady=2)
        self.top_p_scale_val_label = ttk.Label(gen_settings_frame, text=f"{self.top_p_var.get():.2f}", width=4)
        self.top_p_scale_val_label.grid(row=3, column=2, sticky="w", padx=(5,0))

        output_settings_frame = ttk.LabelFrame(settings_container_frame, text="Output File Settings", padding="10")
        output_settings_frame.grid(row=0, column=1, sticky="nsew", padx=5, pady=5)
        output_settings_frame.columnconfigure(1, weight=1)

        ttk.Label(output_settings_frame, text="Output Directory:").grid(row=0, column=0, sticky="w", pady=2)
        self.output_dir_entry = ttk.Entry(output_settings_frame, textvariable=self.output_dir_var)
        self.output_dir_entry.grid(row=0, column=1, sticky="ew", pady=2)
        ttk.Button(output_settings_frame, text="Browse...", command=self.browse_main_output_directory).grid(row=0, column=2, padx=(5,0))

        ttk.Label(output_settings_frame, text="Story/Base Name:").grid(row=1, column=0, sticky="w", pady=2)
        self.story_name_entry = ttk.Entry(output_settings_frame, textvariable=self.story_name_var)
        self.story_name_entry.grid(row=1, column=1, columnspan=2, sticky="ew", pady=2)

        action_buttons_frame = ttk.Frame(self, padding="5")
        action_buttons_frame.grid(row=3, column=0, pady=(10,0), sticky="ew", padx=5)
        self.generate_button = ttk.Button(action_buttons_frame, text="Generate Voice", command=self.start_tts_thread, style="Accent.TButton")
        self.generate_button.pack(side="left", padx=5)
        self.open_folder_button = ttk.Button(action_buttons_frame, text="Open Output Folder", command=self.open_last_output_folder, state="disabled")
        self.open_folder_button.pack(side="left", padx=5)

        progress_frame = ttk.LabelFrame(self, text="Progress", padding="10")
        progress_frame.grid(row=4, column=0, sticky="ewns", padx=5, pady=5)
        progress_frame.columnconfigure(1, weight=1)
        overall_progress_subframe = ttk.Frame(progress_frame)
        overall_progress_subframe.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 5))
        overall_progress_subframe.columnconfigure(1, weight=1)
        ttk.Label(overall_progress_subframe, text="Overall:", width=10).grid(row=0, column=0, sticky="w")
        self.progress_bar_total = ttk.Progressbar(overall_progress_subframe, orient="horizontal", mode="determinate")
        self.progress_bar_total.grid(row=0, column=1, sticky="ew", padx=(0, 10))
        self.total_time_label = ttk.Label(overall_progress_subframe, text="Tổng thời gian: 00:00", style="secondary.TLabel")
        self.total_time_label.grid(row=0, column=2, sticky="e")

        for i in range(NUM_API_KEYS):
            ttk.Label(progress_frame, text=f"API Key {i+1}:", width=10).grid(row=i+1, column=0, sticky="w", pady=1)
            status_label = ttk.Label(progress_frame, text="Idle", foreground=COLOR_NORMAL, anchor="w")
            status_label.grid(row=i+1, column=1, sticky="ew", pady=1, padx=5)
            self.thread_status_labels.append(status_label)

    def set_script_content(self, content):
        """Hàm công khai để main_app có thể điền kịch bản vào."""
        self.main_text_input.delete("1.0", tk.END)
        self.main_text_input.insert("1.0", content)
        self.update_word_count()

    def _on_language_change(self, event=None):
        lang_key = self.languages[self.selected_language.get()]
        new_styles = PREDEFINED_READING_STYLES.get(lang_key, [])
        self.reading_style_combobox['values'] = new_styles
        if new_styles:
            self.reading_style_combobox.current(0)
        else:
            self.reading_style_combobox.set('')
        self.main_app.log_message(f"Đã chuyển ngôn ngữ sang: {self.selected_language.get()}")

    def on_voice_selected(self, event):
        selected_index = event.widget.current()
        if 0 <= selected_index < len(self.voice_name_list):
            voice_name = self.voice_name_list[selected_index]
            self.selected_voice_name.set(voice_name)

    def update_voice_display(self, voice_name_to_set):
        try:
            if voice_name_to_set in self.voice_name_list:
                idx = self.voice_name_list.index(voice_name_to_set)
                self.selected_voice_display.set(self.voice_display_list[idx])
            elif self.voice_display_list:
                self.selected_voice_display.set(self.voice_display_list[0])
                self.selected_voice_name.set(self.voice_name_list[0])
        except (ValueError, IndexError):
            pass

    def import_text_from_file(self):
        file_path = filedialog.askopenfilename(
            title="Chọn file văn bản",
            filetypes=[("Word Document", "*.docx"), ("Text File", "*.txt"), ("All files", "*.*")]
        )
        if not file_path: return
        content = ""
        try:
            if file_path.lower().endswith(".txt"):
                with open(file_path, 'r', encoding='utf-8') as f: content = f.read()
            elif file_path.lower().endswith(".docx"):
                if docx is None:
                    messagebox.showerror("Thiếu thư viện", "Vui lòng cài đặt 'python-docx'.")
                    return
                doc = docx.Document(file_path)
                content = "\n".join([para.text for para in doc.paragraphs])
            else:
                messagebox.showwarning("Định dạng không hỗ trợ", "Chỉ hỗ trợ file .txt và .docx.")
                return
            self.set_script_content(content)
            self.main_app.log_message(f"Đã import thành công từ: {os.path.basename(file_path)}")
        except Exception as e:
            messagebox.showerror("Lỗi đọc file", f"Lỗi: {e}")
            self.main_app.log_message(f"Lỗi import file: {e}")

    def browse_main_output_directory(self):
        directory = filedialog.askdirectory(initialdir=self.output_dir_var.get(), title="Select Output Directory")
        if directory:
            self.output_dir_var.set(directory)
            self.main_app.log_message(f"Output directory set to: {directory}")

    def open_last_output_folder(self):
        if self.last_saved_output_dir and os.path.isdir(self.last_saved_output_dir):
            try:
                if platform.system() == "Windows": os.startfile(self.last_saved_output_dir)
                elif platform.system() == "Darwin": subprocess.Popen(["open", self.last_saved_output_dir])
                else: subprocess.Popen(["xdg-open", self.last_saved_output_dir])
            except Exception as e:
                self.main_app.log_message(f"Lỗi mở thư mục: {e}")
        else:
            self.main_app.log_message("Không có thư mục output để mở.")

    def update_word_count(self, event=None):
        content = self.main_text_input.get("1.0", tk.END).strip()
        char_count = len(content)
        word_count = len(content.split()) if content else 0
        est_chunks = (word_count + self.words_per_chunk_var.get() - 1) // self.words_per_chunk_var.get() if word_count > 0 else 0
        self.main_word_count_label.config(text=f"Ký tự: {char_count} | Từ: {word_count} | Dự kiến chunks: {est_chunks}")

    def update_overall_progress(self, value: int):
        if hasattr(self, 'progress_bar_total') and self.progress_bar_total.winfo_exists():
            self.progress_bar_total["value"] = value

    def update_thread_progress(self, thread_index: int, status_message: str, status_color: str):
        if thread_index < len(self.thread_status_labels) and self.thread_status_labels[thread_index].winfo_exists():
            self.thread_status_labels[thread_index].config(text=status_message, foreground=status_color)

    def start_tts_thread(self):
        text_to_convert = self.main_text_input.get("1.0", tk.END).strip()
        if not text_to_convert:
            messagebox.showerror("Input Error", "Please enter text.")
            return

        active_keys_info_list = self.main_app.get_active_api_keys()
        if not active_keys_info_list:
            messagebox.showerror("API Key Error", "Please set at least one API Key in Settings.")
            self.main_app.notebook.select(self.main_app.settings_tab)
            return
        
        self.main_app.log_message(f"Found {len(active_keys_info_list)} active API Key(s).")
        
        output_dir = self.output_dir_var.get().strip()
        story_base_name = self.story_name_var.get().strip()
        if not output_dir or not os.path.isdir(output_dir) or not story_base_name:
            messagebox.showerror("Output Error", "Invalid output directory or story name.")
            return

        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_story_name = "".join(c if c.isalnum() or c in (' ', '_') else '_' for c in story_base_name).rstrip()
        safe_voice_name = "".join(c if c.isalnum() else '' for c in self.selected_voice_name.get())
        output_filename_only = f"{safe_voice_name}_{safe_story_name}_{timestamp}"
        output_file_path_base = os.path.normpath(os.path.join(output_dir, output_filename_only))

        self.generate_button.config(state="disabled", text="Generating...")
        self.open_folder_button.config(state="disabled")
        self.main_app.root.config(cursor="watch")

        thread = threading.Thread(
            target=self._tts_task_wrapper,
            args=(text_to_convert, output_file_path_base, active_keys_info_list),
            daemon=True
        )
        thread.start()

    def _tts_task_wrapper(self, text_to_convert, output_file_path_base, active_keys_info):
        task_start_time = time.time()
        self.main_app.root.after_idle(lambda: self.total_time_label.config(text="Tổng thời gian: 00:00"))

        try:
            thread_callbacks = [lambda msg, color, i=i: self.main_app.root.after_idle(self.update_thread_progress, i, msg, color) for i in range(len(active_keys_info))]

            success, final_file_path = generate_tts_audio_multithreaded(
                active_api_keys_info=active_keys_info,
                text_to_speak=text_to_convert,
                voice_name=self.selected_voice_name.get(),
                output_file_path_base=output_file_path_base,
                log_callback_ui=self.main_app.log_message,
                progress_callback_ui_total=lambda v: self.main_app.root.after_idle(self.update_overall_progress, v),
                progress_callbacks_ui_thread=thread_callbacks,
                reading_style_prompt=self.reading_style_prompt_var.get(),
                temp_setting=self.temperature_var.get(),
                top_p_setting=self.top_p_var.get(),
                max_words_per_part=self.words_per_chunk_var.get()
            )
            if success and final_file_path:
                self.last_saved_output_dir = os.path.dirname(final_file_path)
        finally:
            def _finalize_ui():
                duration = time.time() - task_start_time
                minutes, seconds = divmod(duration, 60)
                self.generate_button.config(state="normal", text="Generate Voice")
                self.main_app.root.config(cursor="")
                if self.last_saved_output_dir:
                    self.open_folder_button.config(state="normal")
                self.total_time_label.config(text=f"Tổng thời gian: {int(minutes):02d}:{int(seconds):02d}")
            self.main_app.root.after_idle(_finalize_ui)